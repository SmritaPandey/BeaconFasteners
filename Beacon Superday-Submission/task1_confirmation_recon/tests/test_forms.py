"""
Tests for the vendor forms (vendor_forms.py): the PO acknowledgment form and the
supplier onboarding pack.

Each test writes a real form, then plays the VENDOR by filling it in with
openpyxl (SYNTHETIC vendor answers, clearly made up), then reads it back the
way the daily check does. The ERP-backed tests use the real case data if present.
"""
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date, datetime

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import parsers       # noqa: E402
import recon         # noqa: E402
import vendor_forms as vf  # noqa: E402

DATA = os.environ.get("BEACON_DATA", os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))),
                                                  "Beacon Superday-Candidate", "data"))
ERP = os.path.join(DATA, "beacon_erp.db")
AS_OF = date(2026, 5, 17)
VENDOR = {"vendor_id": "V001", "vendor_name": "Synthetic Bar Co.", "country": "USA", "ap_email": "orders@example.com"}


def po_line(line=1, pn="BAR-X", qty=100, price=2.0, required=date(2026, 6, 10)):   # SYNTHETIC PO line
    return dict(po_number="PO-9000000001", po_date=date(2026, 5, 1), vendor_id="V001", vendor_name="Synthetic Bar Co.",
                line_number=line, our_pn=pn, description="Synthetic part %s" % pn, qty=float(qty), price=price,
                required=required)


PO = [po_line(1, "BAR-X", 100, 2.0), po_line(2, "BAR-Y", 50, 7.5)]


class FormCase(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp()
        self.path = os.path.join(self.dir, "ack.xlsx")
        vf.write_ack_form(self.path, PO, VENDOR, today=date(2026, 5, 2))

    def tearDown(self):
        shutil.rmtree(self.dir, ignore_errors=True)

    def fill(self, rows, sign=True):
        """rows: {line_no: dict(resp=, qty=, price=, ccy=, date=, reason=, comment=, own=)} - SYNTHETIC vendor answers."""
        wb = openpyxl.load_workbook(self.path)
        ws = wb["Acknowledgment"]
        col = dict(own=3, resp=9, qty=10, price=11, ccy=12, date=13, reason=14, comment=15)
        for i, pl in enumerate(PO):
            for k, v in rows.get(pl["line_number"], {}).items():
                ws.cell(row=vf.ACK_FIRST_ROW + 1 + i, column=col[k], value=v)
        if sign:
            sr = vf.ACK_FIRST_ROW + len(PO) + 2
            ws.cell(row=sr, column=4, value="Pat Vendor")
            ws.cell(row=sr + 1, column=4, value="pat@example.com")
            ws.cell(row=sr + 2, column=4, value=datetime(2026, 5, 4))
        wb.save(self.path)
        return parsers.parse_file(self.path)

    def reconcile(self, d):
        d["file"] = "ack.xlsx"
        return recon.reconcile(PO, [d], {}, {}, AS_OF)


class TestAckForm(FormCase):
    def test_form_is_locked_and_identified(self):
        ws = openpyxl.load_workbook(self.path)["Acknowledgment"]
        self.assertTrue(ws.protection.sheet)                    # rows can't be deleted, PO cells can't be edited
        self.assertTrue(ws.cell(row=vf.ACK_FIRST_ROW + 1, column=2).protection.locked)    # Beacon PN
        self.assertFalse(ws.cell(row=vf.ACK_FIRST_ROW + 1, column=9).protection.locked)   # Response
        self.assertEqual(len(ws.data_validations.dataValidation), 7)

    def test_clean_acceptance_matches_high_with_no_issues(self):
        d = self.fill({1: dict(resp="Accept as ordered", date=datetime(2026, 6, 10)),
                       2: dict(resp="Accept as ordered", date=datetime(2026, 6, 8))})
        self.assertEqual(d["text_source"], "form")
        self.assertEqual(len(d["lines"]), 2)
        self.assertEqual(d["doc_date"], date(2026, 5, 4))
        res, exc, review, sugg, _ = self.reconcile(d)
        for r in res.values():
            self.assertEqual(r["confidence"], recon.HIGH)
            self.assertEqual([i for i in r["issues"] if recon.ISSUE[i][0] <= 3], [])
        self.assertEqual(review, [])

    def test_declined_line_is_not_called_a_dropped_line(self):
        d = self.fill({1: dict(resp="Cannot supply", comment="Mill allocation - no A286 until Q3"),
                       2: dict(resp="Accept as ordered", date=datetime(2026, 6, 10))})
        self.assertIn(1, d["declined"])
        res = self.reconcile(d)[0]
        r = res[("PO-9000000001", 1)]
        self.assertIn("DECLINED", r["issues"])
        self.assertNotIn("DROPPED", r["issues"])
        self.assertEqual(r["sev"], 1)
        self.assertIn("Mill allocation", " ".join(r["notes"]))
        self.assertIn("cannot supply", recon.action_text(r))

    def test_week_range_and_missing_response_make_line_incomplete(self):
        d = self.fill({1: dict(resp="Accept as ordered", date="KW 20-22"),
                       2: dict(date=datetime(2026, 6, 10))})
        self.assertIn("not a date", d["incomplete"][1])
        self.assertIn("no response", d["incomplete"][2])
        res = self.reconcile(d)[0]
        self.assertIn("FORM_INCOMPLETE", res[("PO-9000000001", 1)]["issues"])
        self.assertIn("FORM_INCOMPLETE", res[("PO-9000000001", 2)]["issues"])

    def test_price_change_without_reason_is_flagged_and_checked(self):
        d = self.fill({1: dict(resp="Accept as ordered", price=2.04, date=datetime(2026, 6, 10)),
                       2: dict(resp="Accept with changes", qty=40, reason="Partial now, balance later",
                               date=datetime(2026, 6, 20))})
        self.assertTrue(any("no reason" in w for w in d["warnings"]))
        self.assertTrue(any("Accept as ordered" in w for w in d["warnings"]))
        res = self.reconcile(d)[0]
        self.assertIn("PRICE_UP", res[("PO-9000000001", 1)]["issues"])
        r2 = res[("PO-9000000001", 2)]
        self.assertIn("QTY_SHORT", r2["issues"])
        self.assertIn("LATE", r2["issues"])
        self.assertIn("Partial now", " ".join(r2["notes"]))

    def test_unsigned_form_warns(self):
        d = self.fill({1: dict(resp="Accept as ordered", date=datetime(2026, 6, 10)),
                       2: dict(resp="Accept as ordered", date=datetime(2026, 6, 10))}, sign=False)
        self.assertTrue(any("not signed" in w for w in d["warnings"]))

    def test_vendor_declared_part_number_goes_to_review_not_crosswalk(self):
        d = self.fill({1: dict(resp="Accept as ordered", own="SYN-100", date=datetime(2026, 6, 10)),
                       2: dict(resp="Accept as ordered", own="SAME", date=datetime(2026, 6, 10))})
        sugg = self.reconcile(d)[3]
        self.assertIn(("V001", "SYN-100"), sugg)
        self.assertEqual(sugg[("V001", "SYN-100")]["review_status"], "PENDING REVIEW")
        self.assertNotIn(("V001", "SAME"), sugg)

    def test_other_excel_file_is_unreadable_not_a_crash(self):
        p = os.path.join(self.dir, "random.xlsx")
        wb = openpyxl.Workbook()
        wb.active["A1"] = "hello"
        wb.save(p)
        self.assertEqual(parsers.parse_file(p)["doc_type"], "UNREADABLE")


@unittest.skipUnless(os.path.exists(ERP), "case data not found (set BEACON_DATA)")
class TestErpAndOnboarding(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.profile = vf.part_profile(ERP)
        cls.po_lines = recon.load_pos(os.path.join(DATA, "open_pos.csv"))
        cls.vendors = recon.load_vendors(os.path.join(DATA, "vendor_master.csv"))
        cls.xw = recon.load_crosswalk(os.path.join(os.path.dirname(HERE), "pn_crosswalk.csv"))
        cls.dir = tempfile.mkdtemp()

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.dir, ignore_errors=True)

    def test_part_profile_from_erp(self):
        p = self.profile[("V004", "HT-SOLN")]
        self.assertGreater(p["lines"], 10)
        self.assertLess(p["lt_planned"], p["lt_quoted"])     # Beacon plans heat treat shorter than Continental quotes

    def test_old_open_balances_match_readout(self):
        b = vf.old_open_balances(ERP)
        self.assertEqual(len(b), 40)
        apex = sum(x["balance_value"] for x in b if x["vendor_id"] == "V001")
        self.assertAlmostEqual(apex, 76706.32, places=1)

    def pack(self, vid, answers, company=None):
        """Write the real pack for vid, then fill it as the vendor (SYNTHETIC answers)."""
        path = os.path.join(self.dir, "%s.xlsx" % vid)
        vf.write_onboarding_pack(path, vid, self.vendors, self.profile, self.po_lines, self.xw, today=AS_OF)
        wb = openpyxl.load_workbook(path)
        c = wb["Company"]
        base = dict(ack_name="Sam Supplier", ack_email="acks@example.com", ar_email="ar@example.com", ack_days="2",
                    form_ok="Yes", cert1="AS9100", cert1_no="C-1", cert1_exp=datetime(2027, 12, 31))
        base.update(company or {})
        for i, (k, *_rest) in enumerate(vf.COMPANY_FIELDS):
            if k in base:
                c.cell(row=3 + i, column=2, value=base[k])
        p = wb["Parts"]
        pns = [p.cell(row=vf.PARTS_FIRST_ROW + 1 + i, column=1).value for i in range(p.max_row - vf.PARTS_FIRST_ROW)]
        for i, pn in enumerate(pns):
            a = answers.get(pn, answers.get("*", {}))
            for col, key in ((9, "own"), (10, "lt"), (13, "sur")):
                if key in a:
                    p.cell(row=vf.PARTS_FIRST_ROW + 1 + i, column=col, value=a[key])
        wb.save(path)
        return vf.read_onboarding_pack(path, today=AS_OF), pns

    def test_onboarding_pack_lists_erp_parts_and_catches_renumbering(self):
        res, pns = self.pack("V002", {"CHB-9472-4": dict(own="APH-441-OS", lt=35, sur="No"),
                                      "*": dict(own="SAME", lt=35, sur="No")})
        self.assertIn("CHB-9472-4", pns)
        self.assertEqual(res["problems"], [])
        m = [x for x in res["mappings"] if x["beacon_pn"] == "CHB-9472-4"]
        self.assertEqual(len(m), 1)
        self.assertIn("APH-4411", m[0]["note"])       # ERP knows the old number

    def test_onboarding_problems_duplicates_and_expired_cert(self):
        res, pns = self.pack("V002", {"*": dict(own="APH-DUP", sur="No")},
                             company=dict(ack_email="not-an-email", cert1_exp=datetime(2025, 1, 1)))
        text = " ".join(res["problems"])
        self.assertIn("standard lead time is missing", text)
        self.assertIn("not a valid email", text)
        self.assertIn("expired", text)
        self.assertIn("APH-DUP", text)
        self.assertEqual(res["mappings"], [])           # ambiguous vendor PN is never mapped

    def test_lead_time_gap_for_planning(self):
        res, _ = self.pack("V004", {"*": dict(own="SAME", lt=28, sur="No")})
        gaps = {g["beacon_pn"]: g for g in res["lead_time_gaps"]}
        self.assertIn("HT-SOLN", gaps)
        self.assertEqual(gaps["HT-SOLN"]["vendor_standard"], 28)

    def test_ack_forms_for_every_open_po(self):
        groups = vf.po_groups(self.po_lines)
        for po, lines in groups.items():
            p = os.path.join(self.dir, "%s.xlsx" % po)
            vf.write_ack_form(p, lines, self.vendors[lines[0]["vendor_id"]], self.xw, self.profile)
            d = vf.read_ack_form(p)       # unfilled: every line must be reported incomplete, none silently dropped
            self.assertEqual(sorted(d["incomplete"]), sorted(pl["line_number"] for pl in lines))
            self.assertEqual(d["lines"], [])


if __name__ == "__main__":
    unittest.main()
