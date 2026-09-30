"""
Tests for 'past and future': what happens on day 2.

    python3 -m unittest discover -s tests -v

Unit tests use SYNTHETIC items. The day-2 end-to-end test uses the real case
files for day 1, then simulates Lisa's Excel edits and one SYNTHETIC revised
vendor PDF (generated here, clearly fake) arriving on day 2.
"""
import os
import shutil
import sys
import tempfile
import unittest
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import recon      # noqa: E402
import store as st  # noqa: E402

DATA = os.environ.get("BEACON_DATA", os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))),
                                                  "Beacon Superday-Candidate", "data"))


def item(code="QUANTITY_SHORT", sev=2, po="PO-1", line=1):
    return dict(po_number=po, line_no=line, vendor_id="V001", issue_code=code, severity=sev, title=code,
                detail="d", action="a", impact=10.0, source_files="f.pdf")


class TestStore(unittest.TestCase):
    def setUp(self):
        self.s = st.Store(os.path.join(tempfile.mkdtemp(), "m.db"))

    def test_new_item_opens(self):
        new, res = self.s.sync_items({"PO-1|1": item()}, "2026-05-17")
        self.assertEqual((new, res), (1, 0))
        self.assertEqual(self.s.items(st.OPEN)[0]["status"], "Open")

    def test_done_stays_done_and_note_kept(self):
        self.s.sync_items({"PO-1|1": item()}, "2026-05-17")
        self.s.set_item("PO-1|1", status=st.DONE, note="Called Apex, balance ships 6/20")
        self.s.sync_items({"PO-1|1": item()}, "2026-05-18")
        i = self.s.items()[0]
        self.assertEqual((i["status"], i["note"]), ("Done", "Called Apex, balance ships 6/20"))

    def test_done_reopens_only_for_a_new_kind_of_problem(self):
        self.s.sync_items({"PO-1|1": item("QUANTITY_SHORT")}, "2026-05-17")
        self.s.set_item("PO-1|1", status=st.DONE)
        self.s.sync_items({"PO-1|1": item("DATE_LATE,QUANTITY_SHORT")}, "2026-05-18")
        i = self.s.items()[0]
        self.assertEqual(i["status"], "Open")
        self.assertIn("new problem", i["status_reason"])

    def test_fixed_by_vendor_auto_resolves(self):
        self.s.sync_items({"PO-1|1": item()}, "2026-05-17")
        new, res = self.s.sync_items({}, "2026-05-18")
        self.assertEqual(res, 1)
        self.assertEqual(self.s.items()[0]["status"], "Resolved")

    def test_first_promise_never_overwritten(self):
        base = dict(po_number="PO-1", line_no=1, confirmed_qty=100, confirmed_price=1.0, currency="USD")
        self.s.record_promise(dict(base, source_file="a.pdf", doc_date=date(2026, 5, 1), promised_date=date(2026, 6, 1)))
        self.s.record_promise(dict(base, source_file="b.pdf", doc_date=date(2026, 5, 9), promised_date=date(2026, 6, 8)))
        self.s.record_promise(dict(base, source_file="b.pdf", doc_date=date(2026, 5, 9), promised_date=date(2026, 6, 8)))  # re-run
        h = self.s.promise_summary()[("PO-1", 1)]
        self.assertEqual((h["first_promise"], h["latest_promise"], h["changes"]), (date(2026, 6, 1), date(2026, 6, 8), 1))

    def test_mapping_review(self):
        self.s.suggest_mapping(dict(vendor_id="V002", vendor_pn="X-1", beacon_pn="CHB-1", description="", mapping_source="t",
                                    mapping_confidence="MEDIUM", first_seen="2026-05-17", last_seen="2026-05-17"))
        self.assertTrue(self.s.crosswalk()[("V002", "X-1")]["review_status"].startswith("PENDING"))
        self.s.review_mapping("V002", "X-1", True, "Lisa")
        self.assertEqual(self.s.crosswalk()[("V002", "X-1")]["review_status"], "APPROVED")


class TestManualAndAI(unittest.TestCase):
    def test_manual_doc(self):
        d = recon.manual_doc(dict(vendor_id="V004", po_number="PO-4500050022", doc_date="2026-05-14",
                                  lines=[dict(line_no=1, vendor_pn="", description="Precip", qty=500, unit_price="",
                                              currency="USD", promise_date="2026-06-11")]))
        self.assertEqual(d["text_source"], "manual")
        self.assertEqual(d["lines"][0]["promise_date"], date(2026, 6, 11))
        self.assertIsNone(d["lines"][0]["unit_price"])

    def test_ai_json_maps_to_doc_shape(self):
        import ai_reader
        data = dict(vendor_name="Ostmark Werkzeug GmbH", po_number="PO-4500050027", doc_type="ACK", doc_date="2026-05-01",
                    vendor_ref="OST-1", notes="", lines=[dict(line_no=2, vendor_pn="OST-CAR-A-100", description="", qty=50,
                                                              unit_price=16.744, currency="eur", promise_date="2026-05-31",
                                                              promise_text="KW 20-22 / 2026")])
        d = ai_reader.to_doc(data, {"V005": {"vendor_name": "Ostmark Werkzeug GmbH"}})
        self.assertEqual((d["vendor_id"], d["po_number"], d["text_source"]), ("V005", "PO-4500050027", "ai"))
        self.assertEqual(d["lines"][0]["currency"], "EUR")
        self.assertIn("KW 20-22", d["lines"][0]["promise_note"])

    def test_ai_off_by_default(self):
        import ai_reader
        os.environ.pop("BEACON_AI_FALLBACK", None)
        self.assertFalse(ai_reader.enabled())


@unittest.skipUnless(os.path.isdir(os.path.join(DATA, "confirmations")), "case data not found (set BEACON_DATA)")
class TestDayTwo(unittest.TestCase):
    """Day 1: real files. Lisa edits the workbook. Day 2: a SYNTHETIC revised ack arrives."""

    def test_day_two(self):
        import fitz
        import openpyxl
        work = tempfile.mkdtemp()
        inbox = os.path.join(work, "inbox")
        os.makedirs(inbox)
        for f in os.listdir(os.path.join(DATA, "confirmations")):   # copy files, not the read-only folder
            shutil.copyfile(os.path.join(DATA, "confirmations", f), os.path.join(inbox, f))
        out = os.path.join(work, "out")
        args = (inbox, os.path.join(DATA, "open_pos.csv"), os.path.join(DATA, "vendor_master.csv"),
                os.path.join(DATA, "beacon_erp.db"), None)

        day1 = recon.run(*args, as_of="2026-05-17", out=out)
        self.assertIn("QTY_SHORT", day1["results"][("PO-4500050007", 1)]["issues"])

        # Lisa works in Excel: marks the dropped line Done with a note, approves Heritage's new PN
        wb = openpyxl.load_workbook(day1["xlsx"])
        ws = wb["Action List"]
        hdr = [c.value for c in ws[4]]
        for row in ws.iter_rows(min_row=5):
            if row[hdr.index("key")].value == "PO-4500050001|2":
                row[hdr.index("Done?")].value = "Yes"
                row[hdr.index("Notes")].value = "Apex: line 2 ships with PO-4500050008, confirmed by phone"
        ws = wb["PN Crosswalk"]
        hdr = [c.value for c in ws[3]]
        for row in ws.iter_rows(min_row=4):
            if row[hdr.index("Vendor PN")].value == "APH-441-OS":
                row[hdr.index("Review status")].value = "APPROVED"
        wb.save(day1["xlsx"])

        # Day 2: SYNTHETIC revised Apex ack for PO-...007 now confirms the full 1,500 pcs
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((50, 60), "\n".join([
            "APEX BAR & TUBE CO.", "ORDER ACKNOWLEDGMENT (REVISED) - SYNTHETIC TEST DOCUMENT",
            "Customer PO: PO-4500050007", "Acknowledgment Date: 05/18/2026", "Apex Order #: APX-978346-R1",
            "Line     Customer PN               Description                            Qty           Unit Price    Promise Date",
            "1        BAR-A286-250              A286 bar stock, 0.250 dia, mill-cert         1,500       $4.7100      06/15/2026"]),
            fontname="cour", fontsize=7)
        doc.save(os.path.join(inbox, "apex_07_revised_SYNTHETIC.pdf"))

        day2 = recon.run(*args, as_of="2026-05-18", out=out)
        s = day2["store"]
        items = {i["item_key"]: i for i in s.items()}
        # vendor fixed it -> closes by itself, with the reason
        self.assertEqual(items["PO-4500050007|1"]["status"], "Resolved")
        # Lisa's decision and note survived
        self.assertEqual(items["PO-4500050001|2"]["status"], "Done")
        self.assertIn("confirmed by phone", items["PO-4500050001|2"]["note"])
        # approved PN is now a normal, high-confidence match; that line is clean -> resolved
        self.assertEqual(day2["results"][("PO-4500050015", 1)]["confidence"], "HIGH")
        self.assertEqual(items["PO-4500050015|1"]["status"], "Resolved")
        # nothing already-seen was re-read as new
        self.assertEqual(day2["stats"]["New documents this run"], 1)
        self.assertEqual(day2["stats"]["Edits read back from last workbook"], 2)


if __name__ == "__main__":
    unittest.main()
