"""
The daily check, shared by the command line (recon.py) and the Desk app (desk_app.py).

    check(store, paths)  ->  reads every stored document, reconciles against
                             today's open PO list, and updates the memory:
                             - to-do items (new / still open / auto-resolved)
                             - promise history (first promise is never overwritten)
                             - new vendor part numbers (pending Lisa's OK)

Both front doors call this, so Excel and the app always agree.
"""
import os
from datetime import date, datetime

import recon
from store import DONE, OPEN


class Paths:
    def __init__(self, open_pos, vendors, erp=None):
        self.open_pos, self.vendors, self.erp = open_pos, vendors, erp


def vendor_history(erp_path):
    """Past performance per vendor from ERP history (Task 2 analysis); {} if unavailable."""
    if not erp_path or not os.path.exists(erp_path):
        return {}
    import vendor_forms
    try:
        s = vendor_forms.erp_build(erp_path)["score"]
    except Exception:
        return {}
    return {vid: dict(on_time=float(r.on_time), meets_promise=float(r.on_time_vs_promise),
                      median_late=float(r.median_days_late_when_late), severe=int(r.severe_lines),
                      promised_late=float(r.pct_promised_late), lt_allowed=float(r.lt_allowed),
                      lt_quoted=float(r.lt_promised), lt_actual=float(r.lt_actual), due=int(r.due_lines))
            for vid, r in s.iterrows()}


def item_key(r=None, exc=None):
    """Stable identity of a to-do item. A PO line keeps ONE item across days,
    so Lisa's note survives even if the details change."""
    if exc is not None:
        return "doc|%s|%s" % (exc["file"], exc["code"])
    return "%s|%d" % (r["po_number"], r["line_number"])


def check(store, paths, reader=None, hist=None, as_of=None):
    reader = reader or (lambda m, vendors, manual: recon.read_document(m["stored_path"], vendors, manual=manual))
    cfg = store.settings()
    po_lines = recon.load_pos(paths.open_pos)
    vendors = recon.load_vendors(paths.vendors)
    fx = recon.load_fx(paths.erp)
    hist = hist if hist is not None else vendor_history(paths.erp)

    open_pos = {p["po_number"] for p in po_lines}
    seen_pos = set(cfg.get("seen_pos", [])) | open_pos
    store.save_setting("seen_pos", sorted(seen_pos))

    docs = []
    for m in store.documents():
        d = reader(m, vendors, store.manual(m["file_hash"]))
        d = dict(d, lines=[dict(l) for l in d["lines"]], warnings=list(d["warnings"]), file=m["file_name"], hash=m["file_hash"])
        docs.append(d)
    # Paperwork for POs that have since closed is history, not a new problem
    live = [d for d in docs if d.get("po_number") in open_pos or d.get("po_number") not in seen_pos]
    for d in docs:
        if not any(d is x for x in live):
            d["status"] = "PO closed - archived"
    active = recon.mark_duplicates(live)
    for d in live:
        if d.get("duplicate_of"):
            d["status"] = "Duplicate of %s - ignored" % d["duplicate_of"]

    if as_of is None:
        newest = max([d["doc_date"] for d in active if d.get("doc_date")] or [date.today()])
        if cfg.get("as_of"):
            as_of = datetime.strptime(cfg["as_of"], "%Y-%m-%d").date()
        elif (date.today() - newest).days > 30:
            as_of = newest   # sample/demo data: don't age everything to today
        else:
            as_of = date.today()

    recon.PRICE_TOL_FX_PCT = float(cfg["price_tol_fx_pct"])
    results, exceptions, review, suggestions, extracted = recon.reconcile(po_lines, active, store.crosswalk(), fx, as_of)
    for s in suggestions.values():
        store.suggest_mapping(s)
    for d in docs:
        store.update_doc_status(d["hash"], d.get("status", "Used"))

    # Promise history: append-only; the first promise is kept forever
    for e in extracted:
        store.record_promise(e)
    ph = store.promise_summary()
    for r in results.values():
        h = ph.get((r["po_number"], r["line_number"]))
        r["first_promise"] = h["first_promise"] if h else None
        r["promise_changes"] = h["changes"] if h else 0
        if h and h["changes"] and r["promise"] and h["first_promise"] and r["promise"] > h["first_promise"]:
            r["notes"].append("Vendor has moved this date %d time(s); first promised %s" % (h["changes"], recon.fmt_d(h["first_promise"])))

    current = {}
    for r in results.values():
        shown = [i for i in r["issues"] if recon.ISSUE[i][0] <= 3]
        if not shown:
            continue
        if shown == ["NO_CONF"] and (as_of - r["po_date"]).days < int(cfg["chase_after_days"]):
            continue  # still inside the normal time for a vendor to acknowledge
        h = hist.get(r["vendor_id"])
        hint = ""
        if h and r["promise"] and h["meets_promise"] < 0.8:
            hint = " | History: %s met %.0f%% of its promised dates (median %.0f days late when late)." % (
                r["vendor_name"], 100 * h["meets_promise"], h["median_late"])
        current[item_key(r)] = dict(
            po_number=r["po_number"], line_no=r["line_number"], vendor_id=r["vendor_id"],
            issue_code=",".join(sorted(set(recon.ISSUE[i][1] for i in shown))), severity=r["sev"],
            title="; ".join(recon.ISSUE[i][2] for i in shown), detail="; ".join(r["notes"]) + hint,
            action=recon.action_text(r), impact=r["impact"], source_files=r["files"])
    for e in exceptions:
        current[item_key(exc=e)] = dict(
            po_number=e["po"], line_no=None, vendor_id=e["vendor"], issue_code=e["code"], severity=e["sev"],
            title=e["issue"], detail=e["detail"], action=e["action"], impact=0.0, source_files=e["file"])
    if not docs:   # nothing received yet: don't flag every PO as "no confirmation" on day zero
        current = {}
    new, resolved = store.sync_items(current, as_of.isoformat())
    store.log_run(as_of.isoformat(), len(docs), new, len(store.items(OPEN)), resolved)
    import vendor_forms
    try:
        backlog = vendor_forms.old_open_balances(paths.erp)
    except Exception:
        backlog = []
    return dict(results=results, exceptions=exceptions, review=review, docs=docs, as_of=as_of, vendors=vendors,
                po_lines=po_lines, hist=hist, new=new, resolved=resolved, extracted=extracted, suggestions=suggestions,
                backlog=backlog)


def item_status(store):
    """{item_key: (status, note, first_seen)} for writing into the Excel workbook."""
    return {i["item_key"]: (i["status"], i["note"] or "", i["first_seen"]) for i in store.items()}


def import_excel_feedback(store, xlsx_path):
    """Read Lisa's edits back from the last workbook she worked in:
    Done? / Notes on the Action List, and review_status on the PN Crosswalk."""
    import openpyxl
    if not xlsx_path or not os.path.exists(xlsx_path):
        return 0
    n = 0
    wb = openpyxl.load_workbook(xlsx_path, data_only=True, read_only=True)
    if "Action List" in wb.sheetnames:
        rows = list(wb["Action List"].iter_rows(values_only=True))
        hdr_i = next((i for i, r in enumerate(rows) if r and "Priority" in r), None)
        if hdr_i is not None:
            h = {v: i for i, v in enumerate(rows[hdr_i]) if v}
            for r in rows[hdr_i + 1:]:
                key = r[h["key"]] if "key" in h else None
                if not key:
                    continue
                done = str(r[h["Done?"]] or "").strip().lower()
                note = r[h["Notes"]] or ""
                if done in ("yes", "y", "x", "done"):
                    store.set_item(key, status=DONE, note=str(note))
                    n += 1
                elif note:
                    store.set_item(key, note=str(note))
                    n += 1
    if "PN Crosswalk" in wb.sheetnames:
        rows = list(wb["PN Crosswalk"].iter_rows(values_only=True))
        hdr_i = next((i for i, r in enumerate(rows) if r and "Vendor PN" in r), None)
        if hdr_i is not None:
            h = {v: i for i, v in enumerate(rows[hdr_i]) if v}
            cw = store.crosswalk()
            for r in rows[hdr_i + 1:]:
                vid, vpn = r[h.get("vendor_id", 0)], r[h["Vendor PN"]]
                status = str(r[h["Review status"]] or "").upper()
                cur = cw.get((vid, vpn))
                if cur and cur["review_status"].startswith("PENDING") and status.startswith(("APPROVED", "REJECTED")):
                    store.review_mapping(vid, vpn, status.startswith("APPROVED"), "Excel", beacon_pn=r[h["Beacon PN"]])
                    n += 1
    return n
