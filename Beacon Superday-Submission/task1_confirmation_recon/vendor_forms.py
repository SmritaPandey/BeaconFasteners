"""
Vendor forms: stop discrepancies at the source instead of reading them out of PDFs.

Two Excel files Beacon sends to vendors. Both are pre-filled from the ERP
history and the open PO list, locked where the vendor must not type, and read
back by the same daily check that reads the PDFs.

1. PO acknowledgment form (one per PO, sent with the PO)
   - Our columns (line, part, qty, price, required date) are locked, and rows
     can't be deleted, so a line can't disappear silently.
   - For EVERY line the vendor picks Accept as ordered / Accept with changes /
     Cannot supply, and gives one delivery date (a real date, not "KW 20-22").
   - Any change to qty or price needs a reason code.
   read_ack_form() turns the returned file into the same normalized document
   the PDF parsers produce, so it drops straight into recon.reconcile() and
   matches on Beacon part numbers at HIGH confidence, with no OCR.
   Lines the vendor declined or left incomplete are reported as such rather
   than as "possible dropped line".

2. Vendor onboarding pack (one per vendor, sent once and refreshed yearly)
   - Company sheet: acknowledgment contact, invoicing currency, acknowledgment
     turnaround, quality certifications with expiry.
   - Parts sheet: every part Beacon has bought from this vendor in the ERP,
     with Beacon's own record of planned / quoted / actual lead time. The vendor
     confirms THEIR part number and standard lead time for each.
   read_onboarding_pack() validates it and returns the vendor profile, the
   vendor-declared part-number mappings (which go to Lisa's approval queue,
   never straight into the crosswalk) and lead-time gaps for planning.

Command line:
  python3 vendor_forms.py ack      --pos open_pos.csv --vendors vendor_master.csv --erp beacon_erp.db --out DIR [--po PO-...]
  python3 vendor_forms.py onboard  --pos open_pos.csv --vendors vendor_master.csv --erp beacon_erp.db --out DIR [--vendor V004]
  python3 vendor_forms.py check    FILE.xlsx          (validate a returned form and print the problems)
"""
import argparse
import json
import os
import re
import sys
from collections import defaultdict
from datetime import date, datetime, timedelta

import xlsxwriter

HERE = os.path.dirname(os.path.abspath(__file__))
META = "_beacon"          # very-hidden sheet that identifies a Beacon form
VERSION = 1

RESPONSES = ["Accept as ordered", "Accept with changes", "Cannot supply"]
REASONS = ["Material / alloy surcharge", "Price list change", "Capacity / schedule", "Material availability",
           "Minimum order / pack size", "Partial now, balance later", "Other (explain in comment)"]
CURRENCIES = ["USD", "EUR"]


# --------------------------------------------------------------------------
# Past: what the ERP knows about each vendor + part
# --------------------------------------------------------------------------

_ERP_CACHE = {}


def erp_build(erp):
    """The Task 2 analysis of the ERP (cleaned receipts, confirmations, past-due
    lines). Cached per file version; None if no ERP file."""
    if not erp or not os.path.exists(erp):
        return None
    key = (os.path.abspath(erp), os.path.getmtime(erp))
    if key not in _ERP_CACHE:
        sys.path.insert(0, os.path.join(HERE, "..", "task2_vendor_readout"))
        from analysis import build
        _ERP_CACHE.clear()
        _ERP_CACHE[key] = build(erp)
    return _ERP_CACHE[key]


def part_profile(erp):
    """{(vendor_id, part_id): dict(lines, last_price, uom, lt_planned, lt_quoted, lt_actual, on_time, last_po)}
    from 8 months of ERP history. Lead times are medians in calendar days from PO date."""
    b = erp_build(erp)
    if b is None:
        return {}
    rec = b["rec"].sort_values("po_dt")
    uom = dict(zip(b["part"]["part_id"], b["part"]["uom"]))
    out = {}
    for (vid, pid), g in rec.groupby(["vendor_id", "part_id"]):
        def med(c):
            s = g[c].dropna()
            return int(round(float(s.median()))) if len(s) else None
        rcv = g[g["on_time"].notna()]
        out[(vid, pid)] = dict(
            lines=int(len(g)), last_price=float(g["unit_price"].iloc[-1]), last_po=g["po_number"].iloc[-1],
            uom=uom.get(pid, ""), lt_planned=med("lt_allowed"), lt_quoted=med("lt_promised"), lt_actual=med("lt_actual"),
            on_time=float(rcv["on_time"].astype(float).mean()) if len(rcv) else None)
    return out


def old_open_balances(erp):
    """ERP PO lines past their required date with qty still open: short-shipped
    balances nobody closed, and lines where nothing arrived. Chase or short-close."""
    b = erp_build(erp)
    if b is None:
        return []
    rows = []
    for _, r in b["past_due"].sort_values("balance_value", ascending=False).iterrows():
        rows.append(dict(vendor_id=r["vendor_id"], po_number=r["po_number"], line_no=int(r["line_no"]),
                         part_id=r["part_id"], ordered=float(r["qty_ordered"]), received=float(r["rec_qty"] or 0),
                         balance=float(r["balance_qty"]), balance_value=float(r["balance_value"]),
                         required=r["required"].date(), days_past_due=int(r["days_past_due"]), kind=r["kind"],
                         last_receipt=r["last_rcv"].date() if str(r["last_rcv"]) != "NaT" else None))
    return rows


# --------------------------------------------------------------------------
# Shared Excel bits
# --------------------------------------------------------------------------

def _formats(wb):
    f = dict(
        title=wb.add_format({"bold": True, "font_size": 16, "font_color": "#1F2A33"}),
        sub=wb.add_format({"font_size": 11, "font_color": "#4E5D6A", "text_wrap": True, "valign": "top"}),
        hdr=wb.add_format({"bold": True, "bg_color": "#1F2A33", "font_color": "white", "text_wrap": True,
                           "valign": "top", "border": 1, "locked": True}),
        hdr_in=wb.add_format({"bold": True, "bg_color": "#D9622B", "font_color": "white", "text_wrap": True,
                              "valign": "top", "border": 1, "locked": True}),
        ro=wb.add_format({"bg_color": "#EEF2F5", "border": 1, "locked": True, "valign": "top", "text_wrap": True}),
        ro_num=wb.add_format({"bg_color": "#EEF2F5", "border": 1, "locked": True, "num_format": "#,##0"}),
        ro_px=wb.add_format({"bg_color": "#EEF2F5", "border": 1, "locked": True, "num_format": "0.0000"}),
        ro_date=wb.add_format({"bg_color": "#EEF2F5", "border": 1, "locked": True, "num_format": "mm/dd/yyyy"}),
        ro_pct=wb.add_format({"bg_color": "#EEF2F5", "border": 1, "locked": True, "num_format": "0%"}),
        inp=wb.add_format({"bg_color": "#FFF4C2", "border": 1, "locked": False, "text_wrap": True, "valign": "top"}),
        inp_num=wb.add_format({"bg_color": "#FFF4C2", "border": 1, "locked": False, "num_format": "#,##0"}),
        inp_px=wb.add_format({"bg_color": "#FFF4C2", "border": 1, "locked": False, "num_format": "0.0000"}),
        inp_date=wb.add_format({"bg_color": "#FFF4C2", "border": 1, "locked": False, "num_format": "mm/dd/yyyy"}),
        label=wb.add_format({"bold": True, "border": 1, "bg_color": "#EEF2F5", "locked": True, "text_wrap": True, "valign": "top"}),
        warn=wb.add_format({"bg_color": "#F8C9A8"}),
        bad=wb.add_format({"font_color": "#B83227", "bold": True}),
    )
    return f


def _meta(wb, payload):
    ws = wb.add_worksheet(META)
    ws.write(0, 0, json.dumps(payload, default=str))
    ws.very_hidden()


def _read_meta(wb):
    if META not in wb.sheetnames:
        return None
    try:
        return json.loads(wb[META]["A1"].value)
    except Exception:
        return None


def _as_date(v):
    """A real date or None. Accepts Excel dates and the usual typed formats."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and 30000 < v < 80000:     # Excel serial typed as a number
        return date(1899, 12, 30) + timedelta(days=int(v))
    s = str(v or "").strip()
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%d.%m.%Y", "%m/%d/%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _as_num(v):
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("$", "").strip())
    except ValueError:
        return None


# --------------------------------------------------------------------------
# 1. PO acknowledgment form
# --------------------------------------------------------------------------

ACK_COLS = ["Line", "Beacon part no.", "Your part no.", "Description", "Ordered qty", "UoM", "PO unit price (USD)",
            "Required at Beacon", "Response *", "Confirmed qty *", "Confirmed unit price *", "Price currency *",
            "Delivery date at Beacon *", "Reason for any change", "Comment"]
ACK_FIRST_ROW = 7     # 0-based row of the first PO line


def write_ack_form(path, po_lines, vendor, crosswalk=None, profile=None, today=None):
    """po_lines: the open-PO rows (recon.load_pos format) of ONE PO.
    crosswalk: {(vendor_id, vendor_pn): row} - approved rows pre-fill 'Your part no.'"""
    today = today or date.today()
    po = po_lines[0]
    vid = po["vendor_id"]
    own = {}
    for (v, vpn), x in (crosswalk or {}).items():
        if v == vid and str(x.get("review_status", "")).startswith("APPROVED"):
            own[x["beacon_pn"]] = vpn
    profile = profile or {}
    wb = xlsxwriter.Workbook(path)
    F = _formats(wb)
    ws = wb.add_worksheet("Acknowledgment")
    ws.write(0, 0, "Beacon Fasteners - Purchase Order Acknowledgment", F["title"])
    ws.merge_range(1, 0, 1, 14, "%s   |   %s (%s)   |   PO date %s" % (po["po_number"], vendor.get("vendor_name", ""), vid,
                                                                       po["po_date"].strftime("%m/%d/%Y")), F["sub"])
    ws.merge_range(2, 0, 3, 14,
                   "Please fill in the YELLOW cells for EVERY line and return this file by reply email. "
                   "Grey cells are our PO and are locked. Choose a Response for each line. If quantity, price or date "
                   "differs from our PO, enter your value and choose a reason. Delivery date = the date the goods "
                   "arrive at Beacon (one date, e.g. 06/15/2026). If you cannot supply a line, choose 'Cannot supply' "
                   "and say why in Comment - please do not delete it.", F["sub"])
    widths = [6, 16, 16, 38, 11, 6, 12, 12, 20, 12, 13, 10, 14, 26, 30]
    for c, w in enumerate(widths):
        ws.set_column(c, c, w)
    ws.set_row(ACK_FIRST_ROW - 1, 32)
    for c, h in enumerate(ACK_COLS):
        ws.write(ACK_FIRST_ROW - 1, c, h, F["hdr_in"] if c >= 8 or c == 2 else F["hdr"])

    meta_lines = []
    last = ACK_FIRST_ROW + len(po_lines) - 1
    for i, pl in enumerate(sorted(po_lines, key=lambda p: p["line_number"])):
        r = ACK_FIRST_ROW + i
        uom = (profile.get((vid, pl["our_pn"])) or {}).get("uom", "")
        ws.write_number(r, 0, pl["line_number"], F["ro_num"])
        ws.write(r, 1, pl["our_pn"], F["ro"])
        ws.write(r, 2, own.get(pl["our_pn"], ""), F["inp"])
        ws.write(r, 3, pl["description"], F["ro"])
        ws.write_number(r, 4, pl["qty"], F["ro_num"])
        ws.write(r, 5, uom, F["ro"])
        ws.write_number(r, 6, pl["price"], F["ro_px"])
        ws.write_datetime(r, 7, datetime.combine(pl["required"], datetime.min.time()), F["ro_date"])
        ws.write_blank(r, 8, None, F["inp"])
        ws.write_number(r, 9, pl["qty"], F["inp_num"])          # pre-filled with our values: edit only if different
        ws.write_number(r, 10, pl["price"], F["inp_px"])
        ws.write(r, 11, "USD", F["inp"])
        ws.write_blank(r, 12, None, F["inp_date"])
        ws.write_blank(r, 13, None, F["inp"])
        ws.write_blank(r, 14, None, F["inp"])
        meta_lines.append(dict(line=pl["line_number"], pn=pl["our_pn"], qty=pl["qty"], price=pl["price"],
                               required=pl["required"].isoformat(), own_pn=own.get(pl["our_pn"], "")))

    rng = lambda c: xlsxwriter.utility.xl_range(ACK_FIRST_ROW, c, last, c)
    ws.data_validation(rng(8), {"validate": "list", "source": RESPONSES, "error_type": "stop",
                                "input_title": "Response", "input_message": "Pick one for every line."})
    ws.data_validation(rng(9), {"validate": "decimal", "criteria": ">", "value": 0, "error_type": "stop",
                                "error_message": "Enter the quantity you will deliver (a number above 0)."})
    ws.data_validation(rng(10), {"validate": "decimal", "criteria": ">", "value": 0, "error_type": "stop",
                                 "error_message": "Enter your unit price (a number above 0)."})
    ws.data_validation(rng(11), {"validate": "list", "source": CURRENCIES, "error_type": "stop"})
    ws.data_validation(rng(12), {"validate": "date", "criteria": "between", "minimum": today - timedelta(days=7),
                                 "maximum": today + timedelta(days=730), "error_type": "stop",
                                 "input_title": "Delivery date", "input_message": "One date, e.g. 06/15/2026.",
                                 "error_message": "Enter one real date (e.g. 06/15/2026). Week ranges such as "
                                                  "'KW 20-22' or 'TBD' cannot be accepted."})
    ws.data_validation(rng(13), {"validate": "list", "source": REASONS, "error_type": "stop"})
    # Highlight anything different from our PO, and blanks still to fill
    for c, ref in ((9, "E"), (10, "G")):
        ws.conditional_format(ACK_FIRST_ROW, c, last, c, {"type": "formula", "criteria": "=%s%d<>%s%d" % (
            xlsxwriter.utility.xl_col_to_name(c), ACK_FIRST_ROW + 1, ref, ACK_FIRST_ROW + 1), "format": F["warn"]})
    ws.conditional_format(ACK_FIRST_ROW, 12, last, 12, {"type": "formula", "criteria": "=AND(M%d<>\"\",M%d>H%d)" % (
        (ACK_FIRST_ROW + 1,) * 3), "format": F["bad"]})

    sig = last + 2
    for k, (lab, fmt) in enumerate((("Acknowledged by (name) *", F["inp"]), ("Email *", F["inp"]),
                                    ("Date *", F["inp_date"]), ("Your order / reference no.", F["inp"]))):
        ws.merge_range(sig + k, 0, sig + k, 2, lab, F["label"])
        ws.merge_range(sig + k, 3, sig + k, 4, "", fmt)
    ws.data_validation(sig + 2, 3, sig + 2, 3, {"validate": "date", "criteria": ">", "value": date(2020, 1, 1)})
    ws.freeze_panes(ACK_FIRST_ROW, 2)
    ws.protect("", {"format_columns": True, "format_rows": True})
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    _meta(wb, dict(form="PO_ACK", version=VERSION, po_number=po["po_number"], vendor_id=vid,
                   generated=today.isoformat(), first_row=ACK_FIRST_ROW, sig_row=sig, lines=meta_lines))
    wb.close()
    return path


def read_ack_form(path):
    """Returned acknowledgment form -> normalized document (same shape as parsers._doc),
    or None if the file is not a Beacon acknowledgment form.

    Extra keys the reconciler understands:
      declined   {line_no: reason}   vendor chose 'Cannot supply'
      incomplete {line_no: problem}  line can't be used (no response, no/invalid date, no qty)
    """
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    meta = _read_meta(wb)
    if not meta or meta.get("form") != "PO_ACK":
        return None
    ws = wb["Acknowledgment"]
    cell = lambda r, c: ws.cell(row=r + 1, column=c + 1).value
    warnings, lines, declined, incomplete = [], [], {}, {}
    for i, m in enumerate(meta["lines"]):
        r = meta["first_row"] + i
        if cell(r, 0) != m["line"] or cell(r, 1) != m["pn"]:
            warnings.append("Row %d was edited or moved; read as line %d (%s) from the original form." % (r + 1, m["line"], m["pn"]))
        resp = str(cell(r, 8) or "").strip()
        own_pn = str(cell(r, 2) or "").strip()
        qty, price = _as_num(cell(r, 9)), _as_num(cell(r, 10))
        ccy = str(cell(r, 11) or "USD").strip().upper()
        raw_date = cell(r, 12)
        prom = _as_date(raw_date)
        reason, comment = str(cell(r, 13) or "").strip(), str(cell(r, 14) or "").strip()
        tag = "Line %d (%s)" % (m["line"], m["pn"])
        if resp == "Cannot supply":
            declined[m["line"]] = "; ".join(x for x in (reason, comment) if x) or "no reason given"
            continue
        problems = []
        if resp not in RESPONSES:
            problems.append("no response chosen")
        if qty is None or qty <= 0:
            problems.append("no confirmed quantity")
        if prom is None:
            problems.append("delivery date %s" % ("'%s' is not a date" % raw_date if raw_date not in (None, "") else "missing"))
        if problems:
            incomplete[m["line"]] = ", ".join(problems)
            continue
        changed = [k for k, a, b in (("qty", qty, m["qty"]), ("price", price, m["price"])) if a is not None and abs(a - b) > 1e-9]
        if ccy != "USD":
            changed.append("currency")
        if changed and not reason:
            warnings.append("%s: %s changed with no reason given." % (tag, " and ".join(changed)))
        if changed and resp == "Accept as ordered":
            warnings.append("%s: marked 'Accept as ordered' but %s differs from the PO." % (tag, " and ".join(changed)))
        lines.append(dict(line_no=m["line"], vendor_pn=m["pn"], own_pn=own_pn or None, description=None, qty=qty,
                          uom=None, unit_price=price, currency=ccy if ccy in CURRENCIES else "USD", promise_date=prom,
                          promise_note=None, vendor_reason="; ".join(x for x in (reason, comment) if x) or None))
    sr = meta["sig_row"]
    who, email, signed, ref = cell(sr, 3), cell(sr + 1, 3), _as_date(cell(sr + 2, 3)), cell(sr + 3, 3)
    if not who or not email:
        warnings.append("Acknowledgment not signed (name/email missing).")
    doc_date = signed or datetime.fromtimestamp(os.path.getmtime(path)).date()
    return dict(vendor_id=meta["vendor_id"], template="Beacon acknowledgment form v%d" % meta["version"],
                doc_type="ACK", po_number=meta["po_number"], doc_date=doc_date, vendor_ref=str(ref) if ref else None,
                part_of=None, supersedes_prior=False, lines=lines, warnings=warnings, declined=declined,
                incomplete=incomplete, text_source="form", signed_by=who, signed_email=email,
                raw_text="Beacon acknowledgment form for %s, signed by %s <%s> on %s" % (
                    meta["po_number"], who or "?", email or "?", doc_date))


# --------------------------------------------------------------------------
# 2. Vendor onboarding pack
# --------------------------------------------------------------------------

COMPANY_FIELDS = [
    # key, label, required, kind
    ("legal_name", "Legal company name", True, "text"),
    ("country", "Country", True, "text"),
    ("currency", "Currency you invoice Beacon in", True, "currency"),
    ("ack_name", "Order acknowledgment contact - name", True, "text"),
    ("ack_email", "Order acknowledgment contact - email", True, "email"),
    ("ar_email", "Accounts receivable email", True, "email"),
    ("quality_email", "Quality contact email", False, "email"),
    ("ack_days", "We acknowledge Beacon POs within (business days)", True, "ackdays"),
    ("form_ok", "We will acknowledge using Beacon's acknowledgment form", True, "yesno"),
    ("cert1", "Quality certification 1", True, "cert"),
    ("cert1_no", "Certificate number", False, "text"),
    ("cert1_exp", "Certificate expiry date", False, "date"),
    ("cert2", "Quality certification 2 (e.g. Nadcap for heat treat / plating)", False, "cert"),
    ("cert2_no", "Certificate number", False, "text"),
    ("cert2_exp", "Certificate expiry date", False, "date"),
]
CERTS = ["AS9100", "ISO 9001", "Nadcap - heat treating", "Nadcap - chemical processing", "Other", "None"]
PART_COLS = ["Beacon part no.", "Description", "UoM", "Beacon POs (8 mo)", "Last PO price (USD)",
             "Beacon planned lead time (days)", "Your quoted lead time (days, median)", "Actual delivered (days, median)",
             "Your part no. *", "Your standard lead time (calendar days) *", "Minimum order qty", "Price basis *",
             "Surcharges apply? *", "Comment"]
PRICE_BASIS = ["per each", "per lb", "per 100", "per 1000", "per lot"]
SURCHARGE = ["No", "Yes - alloy / material surcharge", "Yes - other (explain)"]
PARTS_FIRST_ROW = 4


def vendor_parts(vendor_id, profile, po_lines):
    """Every part Beacon has bought (ERP) or has open (PO list) with this vendor."""
    parts = {}
    for (vid, pid), p in profile.items():
        if vid == vendor_id:
            parts[pid] = dict(p)
    for pl in po_lines:
        if pl["vendor_id"] == vendor_id:
            parts.setdefault(pl["our_pn"], dict(lines=0, last_price=pl["price"], uom="", lt_planned=None,
                                                lt_quoted=None, lt_actual=None))
            parts[pl["our_pn"]].setdefault("description", pl["description"])
    return parts


def write_onboarding_pack(path, vendor_id, vendors, profile, po_lines, crosswalk=None, descriptions=None, today=None):
    today = today or date.today()
    v = vendors.get(vendor_id, {})
    parts = vendor_parts(vendor_id, profile, po_lines)
    descriptions = descriptions or {}
    own = {}
    for (vid, vpn), x in (crosswalk or {}).items():
        if vid == vendor_id and str(x.get("review_status", "")).startswith("APPROVED"):
            own.setdefault(x["beacon_pn"], []).append(vpn)
    wb = xlsxwriter.Workbook(path)
    F = _formats(wb)

    ws = wb.add_worksheet("Start here")
    ws.set_column(0, 0, 110)
    ws.set_landscape()
    ws.fit_to_pages(1, 1)
    ws.write(0, 0, "Beacon Fasteners - Supplier information pack", F["title"])
    for i, t in enumerate([
        "Prepared for %s (Beacon vendor ID %s) on %s." % (v.get("vendor_name", vendor_id), vendor_id, today.strftime("%m/%d/%Y")),
        "Why: every order acknowledgment we receive is checked line by line against our PO. When your part numbers, "
        "lead times and contacts are on file, your acknowledgments match automatically and we only contact you when "
        "something really differs.",
        "1. 'Company' tab: fill the yellow cells (* = required).",
        "2. 'Parts' tab: one row per part we buy from you. Grey columns are from Beacon's purchasing records (last 8 months). "
        "Tell us YOUR part number for each (write SAME if you use ours) and your standard lead time in calendar days "
        "from PO to delivery at Beacon.",
        "3. Save and return this file by email. Please don't rename tabs or delete rows - we read this file automatically.",
    ], 2):
        ws.write(i, 0, t, F["sub"])
        ws.set_row(i, 32)

    ws = wb.add_worksheet("Company")
    ws.set_column(0, 0, 58)
    ws.set_column(1, 1, 42)
    ws.write(0, 0, "Company and contacts", F["title"])
    prefill = dict(legal_name=v.get("vendor_name", ""), country=v.get("country", ""),
                   currency="EUR" if v.get("country", "").lower() in ("germany", "de", "deutschland") else "USD",
                   ack_email=v.get("ap_email", ""))
    for i, (k, lab, req, kind) in enumerate(COMPANY_FIELDS):
        r = 2 + i
        ws.write(r, 0, lab + (" *" if req else ""), F["label"])
        val = prefill.get(k, "")
        ws.write(r, 1, val, F["inp_date"] if kind == "date" else F["inp"])
        dv = {"currency": {"validate": "list", "source": CURRENCIES},
              "ackdays": {"validate": "list", "source": ["1", "2", "3", "5"]},
              "yesno": {"validate": "list", "source": ["Yes", "No"]},
              "cert": {"validate": "list", "source": CERTS},
              "date": {"validate": "date", "criteria": ">", "value": date(2020, 1, 1),
                       "error_message": "Enter a date, e.g. 12/31/2027."},
              "email": {"validate": "custom", "value": '=ISNUMBER(FIND("@",B%d))' % (r + 1),
                        "error_message": "Enter an email address."}}.get(kind)
        if dv:
            ws.data_validation(r, 1, r, 1, dict(dv, error_type="stop"))
    ws.protect("", {"format_columns": True})
    ws.fit_to_pages(1, 1)

    ws = wb.add_worksheet("Parts")
    ws.write(0, 0, "Parts Beacon buys from you", F["title"])
    ws.merge_range(1, 0, 1, 13, "Grey = Beacon's records (8 months). Yellow = please fill. Lead times are calendar days from "
                                "PO date to delivery at Beacon (medians).", F["sub"])
    ws.set_landscape()
    ws.fit_to_pages(1, 0)
    widths = [16, 38, 6, 10, 12, 13, 14, 14, 18, 16, 12, 12, 22, 30]
    for c, w in enumerate(widths):
        ws.set_column(c, c, w)
    ws.set_row(PARTS_FIRST_ROW - 1, 45)
    for c, h in enumerate(PART_COLS):
        ws.write(PARTS_FIRST_ROW - 1, c, h, F["hdr_in"] if c >= 8 else F["hdr"])
    meta_parts = []
    pids = sorted(parts)
    for i, pid in enumerate(pids):
        p, r = parts[pid], PARTS_FIRST_ROW + i
        ws.write(r, 0, pid, F["ro"])
        ws.write(r, 1, descriptions.get(pid) or p.get("description", ""), F["ro"])
        ws.write(r, 2, p.get("uom", ""), F["ro"])
        ws.write_number(r, 3, p.get("lines", 0), F["ro_num"])
        ws.write_number(r, 4, p.get("last_price") or 0, F["ro_px"])
        for c, k in ((5, "lt_planned"), (6, "lt_quoted"), (7, "lt_actual")):
            ws.write(r, c, p.get(k) if p.get(k) is not None else "", F["ro_num"])
        known = own.get(pid, [])
        ws.write(r, 8, known[-1] if len(known) == 1 else "", F["inp"])
        ws.write_blank(r, 9, None, F["inp_num"])
        ws.write_blank(r, 10, None, F["inp_num"])
        ws.write(r, 11, "per lb" if p.get("uom", "").upper() == "LB" else "per each", F["inp"])
        ws.write_blank(r, 12, None, F["inp"])
        ws.write_blank(r, 13, None, F["inp"])
        meta_parts.append(dict(pn=pid, known_own=known, lt_planned=p.get("lt_planned"), lt_quoted=p.get("lt_quoted"),
                               lt_actual=p.get("lt_actual")))
    if pids:
        last = PARTS_FIRST_ROW + len(pids) - 1
        rng = lambda c: xlsxwriter.utility.xl_range(PARTS_FIRST_ROW, c, last, c)
        ws.data_validation(rng(9), {"validate": "integer", "criteria": "between", "minimum": 1, "maximum": 365,
                                    "error_type": "stop", "error_message": "Whole calendar days, 1 to 365."})
        ws.data_validation(rng(10), {"validate": "decimal", "criteria": ">=", "value": 0, "error_type": "stop"})
        ws.data_validation(rng(11), {"validate": "list", "source": PRICE_BASIS, "error_type": "stop"})
        ws.data_validation(rng(12), {"validate": "list", "source": SURCHARGE, "error_type": "stop"})
    ws.freeze_panes(PARTS_FIRST_ROW, 1)
    ws.protect("", {"format_columns": True, "format_rows": True})
    _meta(wb, dict(form="ONBOARDING", version=VERSION, vendor_id=vendor_id, generated=today.isoformat(),
                   parts=meta_parts))
    wb.close()
    return path


def read_onboarding_pack(path, today=None):
    """-> dict(vendor_id, company, parts, problems, mappings, lead_time_gaps) or None if not a Beacon pack.
    problems: things to send back to the vendor. mappings: vendor-declared part numbers for Lisa's approval.
    lead_time_gaps: parts where the vendor's standard lead time is well above what Beacon plans."""
    import openpyxl
    today = today or date.today()
    wb = openpyxl.load_workbook(path, data_only=True)
    meta = _read_meta(wb)
    if not meta or meta.get("form") != "ONBOARDING":
        return None
    vid = meta["vendor_id"]
    problems, company = [], {}
    ws = wb["Company"]
    for i, (k, lab, req, kind) in enumerate(COMPANY_FIELDS):
        v = ws.cell(row=3 + i, column=2).value
        v = _as_date(v) if kind == "date" else (str(v).strip() if v not in (None, "") else "")
        company[k] = v
        if req and not v:
            problems.append("Company: '%s' is required." % lab)
        if kind == "email" and v and not re.match(r"[^@\s]+@[^@\s]+\.[^@\s]+$", v):
            problems.append("Company: '%s' is not a valid email (%s)." % (lab, v))
        if kind == "date" and v and v < today:
            problems.append("Company: %s certificate expired on %s - please send the renewed one." % (
                company.get(k.replace("_exp", "")) or "quality", v.strftime("%m/%d/%Y")))
    for n in ("1", "2"):
        if company.get("cert" + n) not in ("", "None", None) and not company.get("cert%s_exp" % n):
            problems.append("Company: expiry date missing for %s." % company["cert" + n])

    ws = wb["Parts"]
    parts, mappings, gaps, seen_own = [], [], [], defaultdict(list)
    for i, m in enumerate(meta["parts"]):
        r = PARTS_FIRST_ROW + 1 + i
        val = lambda c: ws.cell(row=r, column=c + 1).value
        pn = m["pn"]
        own = str(val(8) or "").strip()
        lt = _as_num(val(9))
        row = dict(beacon_pn=pn, own_pn=own, std_lead_time=int(lt) if lt else None, moq=_as_num(val(10)),
                   price_basis=val(11), surcharge=val(12), comment=val(13),
                   lt_planned=m["lt_planned"], lt_quoted=m["lt_quoted"], lt_actual=m["lt_actual"])
        parts.append(row)
        if not own:
            problems.append("Parts: %s - your part number is missing (write SAME if you use ours)." % pn)
        if not lt or lt <= 0:
            problems.append("Parts: %s - standard lead time is missing." % pn)
        if not val(12):
            problems.append("Parts: %s - please say whether surcharges apply." % pn)
        own_pn = pn if own.upper() == "SAME" else own
        if own_pn:
            seen_own[own_pn].append(pn)
            if own_pn != pn:
                note = ""
                if m["known_own"] and own_pn not in m["known_own"]:
                    note = "Changed: our records have %s for this part - vendor renumbered?" % ", ".join(m["known_own"])
                if own_pn not in m["known_own"]:
                    mappings.append(dict(vendor_id=vid, vendor_pn=own_pn, beacon_pn=pn, note=note))
        if lt and m["lt_planned"] and lt > m["lt_planned"] + 5 and lt > 1.2 * m["lt_planned"]:
            gaps.append(dict(beacon_pn=pn, planned=m["lt_planned"], vendor_standard=int(lt), quoted=m["lt_quoted"],
                             actual=m["lt_actual"], gap=int(lt) - m["lt_planned"]))
    for own_pn, pns in seen_own.items():
        if len(pns) > 1:
            problems.append("Parts: your part number %s is given for %d different Beacon parts (%s)." % (
                own_pn, len(pns), ", ".join(pns)))
            mappings = [x for x in mappings if x["vendor_pn"] != own_pn]   # ambiguous: never map it
    return dict(vendor_id=vid, company=company, parts=parts, problems=problems, mappings=mappings,
                lead_time_gaps=gaps, generated=meta.get("generated"))


def onboarding_reply_email(result, vendor_name, buyer="Lisa"):
    if not result["problems"]:
        return "Thank you - your supplier information is complete and on file.\n\n%s\nBeacon Fasteners Purchasing" % buyer
    return ("Hello,\n\nThank you for returning the supplier information pack for %s. A few items still need your input:\n\n%s\n\n"
            "Please update the same file and send it back.\n\nThank you,\n%s\nBeacon Fasteners Purchasing") % (
        vendor_name, "\n".join("  - " + p for p in result["problems"]), buyer)


# --------------------------------------------------------------------------
# Batch helpers (CLI and the Desk app)
# --------------------------------------------------------------------------

def po_groups(po_lines):
    g = defaultdict(list)
    for pl in po_lines:
        g[pl["po_number"]].append(pl)
    return g


def descriptions_from_erp(erp):
    b = erp_build(erp)
    return dict(zip(b["part"]["part_id"], b["part"]["description"])) if b is not None else {}


def _load_inputs(a):
    import recon
    po_lines = recon.load_pos(a.pos)
    vendors = recon.load_vendors(a.vendors)
    xw = recon.load_crosswalk(a.crosswalk or os.path.join(HERE, "pn_crosswalk.csv"))
    return po_lines, vendors, xw


def main(argv=None):
    ap = argparse.ArgumentParser(description="Beacon vendor forms: PO acknowledgment forms and onboarding packs.")
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("ack", "onboard"):
        p = sub.add_parser(name)
        p.add_argument("--pos", required=True)
        p.add_argument("--vendors", required=True)
        p.add_argument("--erp")
        p.add_argument("--crosswalk")
        p.add_argument("--out", default="vendor_forms")
        p.add_argument("--po" if name == "ack" else "--vendor")
    c = sub.add_parser("check")
    c.add_argument("file")
    a = ap.parse_args(argv)

    if a.cmd == "check":
        res = read_ack_form(a.file) or read_onboarding_pack(a.file)
        if res is None:
            print("Not a Beacon form.")
            return 1
        if "declined" in res:
            print("PO %s from %s: %d usable line(s)" % (res["po_number"], res["vendor_id"], len(res["lines"])))
            for k, v in res["declined"].items():
                print("  DECLINED line %d: %s" % (k, v))
            for k, v in res["incomplete"].items():
                print("  INCOMPLETE line %d: %s" % (k, v))
            for w in res["warnings"]:
                print("  note: %s" % w)
        else:
            print("Onboarding pack for %s: %d parts, %d problem(s)" % (res["vendor_id"], len(res["parts"]), len(res["problems"])))
            for p in res["problems"]:
                print("  PROBLEM " + p)
            for m in res["mappings"]:
                print("  NEW PN (needs Lisa's OK) %s = %s %s" % (m["vendor_pn"], m["beacon_pn"], m["note"]))
            for g in res["lead_time_gaps"]:
                print("  LEAD TIME %s: Beacon plans %d days, vendor standard %d" % (g["beacon_pn"], g["planned"], g["vendor_standard"]))
        return 0

    po_lines, vendors, xw = _load_inputs(a)
    profile = part_profile(a.erp)
    os.makedirs(a.out, exist_ok=True)
    if a.cmd == "ack":
        groups = po_groups(po_lines)
        for po in ([a.po] if a.po else sorted(groups)):
            f = write_ack_form(os.path.join(a.out, "Beacon_Acknowledgment_%s.xlsx" % po), groups[po],
                               vendors.get(groups[po][0]["vendor_id"], {}), xw, profile)
            print("->", f)
    else:
        desc = descriptions_from_erp(a.erp)
        vids = [a.vendor] if a.vendor else sorted(vendors)
        for vid in vids:
            f = write_onboarding_pack(os.path.join(a.out, "Beacon_Supplier_Pack_%s_%s.xlsx" % (
                vid, re.sub(r"[^A-Za-z0-9]+", "_", vendors.get(vid, {}).get("vendor_name", vid)).strip("_"))),
                vid, vendors, profile, po_lines, xw, desc)
            print("->", f)
    return 0


if __name__ == "__main__":
    sys.exit(main())
